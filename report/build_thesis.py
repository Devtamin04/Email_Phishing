"""Builds report/BaoCao_PhishLens_DeTai.docx using BaoCao_VNSearch_DETai2.docx as the template
(cover page, styles, TOC field are reused; the old body is replaced)."""
import copy
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from thesis_content import APPENDIX, BLOCKS, REFERENCES  # noqa: E402

TEMPLATE = os.path.join(HERE, "template", "mau_bao_cao.docx")  # copy of BaoCao_VNSearch_DETai2.docx
OUT = os.path.join(HERE, "BaoCao_PhishLens_DeTai.docx")
IMG = os.path.join(HERE, "img")
FONT, SZ = "Times New Roman", 26          # half-points (13 pt)

COVER = {  # paragraph index -> new text (None = keep)
    6: "[TÊN MÔN HỌC – MÃ MÔN HỌC]",
    9: "Mã đề tài: [..]",
    10: "Tên đề tài: Phát hiện email lừa đảo thế hệ LLM –",
    11: "tái hiện KD-BiLSTM và xây dựng hệ thống PhishLens Edu",
    22: "Thành phố Hồ Chí Minh, tháng 10 năm 2026",
}

doc = Document(TEMPLATE)
body = doc.element.body


# ------------------------------------------------------------------ cover: replace text keep format
def set_text(p, text):
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    # keep the first run that has text (drawings live in their own runs on the cover)
    first = next((r for r in runs if r.text), runs[0])
    first.text = text
    for r in runs:
        if r is not first and r.text:
            r.text = ""


for idx, text in COVER.items():
    set_text(doc.paragraphs[idx], text)

# ------------------------------------------------------------------ drop old body after the TOC
children = list(body.iterchildren())
sdt_i = next(i for i, el in enumerate(children) if el.tag == qn("w:sdt"))
for el in children[sdt_i + 1:]:
    if el.tag != qn("w:sectPr"):
        body.remove(el)

# rebuild TOC content: heading + a fresh TOC field (Word fills it when fields are updated)
sdt = children[sdt_i]
content = sdt.find(qn("w:sdtContent"))
paras = content.findall(qn("w:p"))
for p in paras[1:]:
    content.remove(p)


def _r(text=None, bold=False, italic=False, size=SZ, color=None, vert=None):
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        f.set(qn(a), FONT)
    rpr.append(f)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rpr.append(c)
    for tag in ("w:sz", "w:szCs"):
        s = OxmlElement(tag)
        s.set(qn("w:val"), str(size))
        rpr.append(s)
    if vert:
        v = OxmlElement("w:vertAlign")
        v.set(qn("w:val"), vert)
        rpr.append(v)
    r.append(rpr)
    if text is not None:
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
    return r


def _fld(kind, instr=None):
    r = _r()
    if kind == "instr":
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        r.append(it)
    else:
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        r.append(fc)
    return r


toc_p = OxmlElement("w:p")
ppr = OxmlElement("w:pPr")
ps = OxmlElement("w:pStyle")
ps.set(qn("w:val"), "TOC1")
ppr.append(ps)
toc_p.append(ppr)
for el in (_fld("begin"), _fld("instr", ' TOC \\o "1-3" \\h \\z \\u '), _fld("separate"),
           _r("Nhấn chuột phải vào đây → Update Field (hoặc F9) để cập nhật mục lục.", italic=True),
           _fld("end")):
    toc_p.append(el)
content.append(toc_p)

# ask Word to refresh fields (TOC, page numbers) when the file is opened
settings = doc.settings.element
uf = OxmlElement("w:updateFields")
uf.set(qn("w:val"), "true")
settings.append(uf)


# ------------------------------------------------------------------ paragraph helpers
def _ppr(p, align="both", first_line=567, before=0, after=0, line=360, keep_next=False,
         page_break=False, tabs=None):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line / 240
    if first_line:
        pf.first_line_indent = Cm(first_line / 567)
    pf.alignment = {"both": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                    "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    pf.keep_with_next = keep_next
    pf.page_break_before = page_break
    if tabs:
        for kind, pos in tabs:
            pf.tab_stops.add_tab_stop(Cm(pos), {"center": 1, "right": 2}[kind])


def _run(p, text, bold=False, italic=False, size=13, color=None, vert=None):
    r = p.add_run(text)
    r.font.name, r.font.size, r.bold, r.italic = FONT, Pt(size), bold, italic
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if vert == "sub":
        r.font.subscript = True
    elif vert == "sup":
        r.font.superscript = True
    return r


def rich(p, text, size=13):
    """**bold** and _{sub} / ^{sup} markup."""
    for part in re.split(r"(\*\*.+?\*\*|_\{[^}]*\}|\^\{[^}]*\})", text):
        if not part:
            continue
        if part.startswith("**"):
            _run(p, part[2:-2], bold=True, size=size)
        elif part.startswith("_{"):
            _run(p, part[2:-1], size=size, vert="sub")
        elif part.startswith("^{"):
            _run(p, part[2:-1], size=size, vert="sup")
        else:
            _run(p, part, size=size)


chap, sec, eq_n, tab_n, fig_n = 0, 0, 0, 0, 0
LABEL = [None]  # None -> chapter number; "PL" for the appendix


def h1(title, numbered=True, label=None):
    global chap, sec, eq_n, tab_n, fig_n
    p = doc.add_paragraph(style="Heading 1")
    LABEL[0] = label
    if label:
        sec = tab_n = fig_n = eq_n = 0
        text = title.upper()
    elif numbered:
        chap += 1
        sec = eq_n = tab_n = fig_n = 0
        text = f"CHƯƠNG {chap}: {title.upper()}"
    else:
        text = title.upper()
    _ppr(p, "center", first_line=0, after=6, keep_next=True, page_break=True)
    _run(p, text, bold=True)


def h2(title):
    global sec
    sec += 1
    p = doc.add_paragraph(style="Heading 2")
    _ppr(p, "left", first_line=567, before=6, keep_next=True)
    _run(p, f"{LABEL[0] or chap}.{sec}. {title}", bold=True)


def para(text):
    p = doc.add_paragraph()
    _ppr(p)
    rich(p, text)


def equation(text):
    global eq_n
    eq_n += 1
    p = doc.add_paragraph()
    _ppr(p, "left", first_line=0, before=3, after=3, tabs=[("center", 8.0), ("right", 16.0)])
    _run(p, "\t")
    for part in re.split(r"(_\{[^}]*\}|\^\{[^}]*\})", text):
        if part.startswith("_{"):
            _run(p, part[2:-1], italic=True, vert="sub")
        elif part.startswith("^{"):
            _run(p, part[2:-1], italic=True, vert="sup")
        elif part:
            _run(p, part, italic=True)
    _run(p, f"\t({chap}.{eq_n})")


def code(text):
    lines = text.strip("\n").split("\n")
    for k, line in enumerate(lines):
        p = doc.add_paragraph()
        _ppr(p, "left", first_line=0, before=4 if k == 0 else 0, after=4 if k == len(lines) - 1 else 0,
             line=240, keep_next=k < len(lines) - 1)
        p.paragraph_format.left_indent = Cm(0.6)
        r = p.add_run(line or " ")
        r.font.name, r.font.size = "Consolas", Pt(10.5)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        pPr.append(shd)


def caption(kind, n, text):
    p = doc.add_paragraph()
    _ppr(p, "center", first_line=0, before=3, after=6)
    _run(p, f"{kind} {LABEL[0] or chap}.{n}. ", bold=True)
    _run(p, text, italic=True)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(cap, rows, widths, size=12):
    global tab_n
    tab_n += 1
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = doc.styles["Table Grid"]
    t.alignment = 1
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths[j])
            p = c.paragraphs[0]
            _ppr(p, "center" if (i == 0 or j > 0 and len(str(val)) < 12) else "left", first_line=0,
                 before=2, after=2, line=276)
            _run(p, str(val), bold=i == 0, size=size, color="FFFFFF" if i == 0 else None)
            if i == 0:
                _shade(c, "1F497D")
        if i == 0:  # repeat header row on page breaks
            trPr = t.rows[0]._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            trPr.append(th)
    caption("Bảng", tab_n, cap)


def figure(name, cap, width):
    global fig_n
    fig_n += 1
    p = doc.add_paragraph()
    _ppr(p, "center", first_line=0, before=6, after=0, line=240, keep_next=True)
    p.add_run().add_picture(os.path.join(IMG, name), width=Cm(width))
    caption("Hình", fig_n, cap)


# ------------------------------------------------------------------ body
def render(blocks):
  for b in blocks:
    kind = b[0]
    if kind == "h1":
        h1(b[1])
    elif kind == "appendix":
        h1(b[1], label="PL")
    elif kind == "h2":
        h2(b[1])
    elif kind == "p":
        para(b[1])
    elif kind == "eq":
        equation(b[1])
    elif kind == "tab":
        table(b[1], b[2], b[3], b[4])
    elif kind == "fig":
        figure(b[1], b[2], b[3])
    elif kind == "code":
        code(b[1])


render(BLOCKS)

h1("TÀI LIỆU THAM KHẢO", numbered=False)
for i, ref in enumerate(REFERENCES, 1):
    p = doc.add_paragraph()
    _ppr(p, "both", first_line=0, after=4)
    p.paragraph_format.left_indent = Cm(0.9)
    p.paragraph_format.first_line_indent = Cm(-0.9)
    _run(p, f"[{i}]\t{ref}")

render(APPENDIX)

# ------------------------------------------------------------------ page numbers (content section)
sec_last = doc.sections[-1]
sec_last.footer.is_linked_to_previous = False
fp = sec_last.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for el in (_fld("begin"), _fld("instr", " PAGE "), _fld("separate"), _r("1"), _fld("end")):
    fp._p.append(el)
sectPr = sec_last._sectPr
pg = sectPr.find(qn("w:pgNumType"))
if pg is None:
    pg = OxmlElement("w:pgNumType")
    sectPr.append(pg)
pg.set(qn("w:start"), "1")
doc.sections[0].footer.is_linked_to_previous = False

doc.save(OUT)
print("saved", OUT)
