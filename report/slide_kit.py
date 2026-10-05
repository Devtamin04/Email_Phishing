"""Shared slide helpers & palette for build_pptx.py and build_slides.py"""
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C  # noqa: E402

INK, INK2, MUTED = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x89, 0x87, 0x81)
ACCENT, SOFT = RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xCD, 0xE2, 0xFB)
PAGE, CARD, GRID = RGBColor(0xF9, 0xF9, 0xF7), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xE1, 0xE0, 0xD9)
SERIES = [RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x1B, 0xAF, 0x7A),
          RGBColor(0xED, 0xA1, 0x00)]
STATUS = {"good": RGBColor(0x0C, 0xA3, 0x0C), "warning": RGBColor(0xFA, 0xB2, 0x19),
          "serious": RGBColor(0xEC, 0x83, 0x5A), "critical": RGBColor(0xD0, 0x3B, 0x3B)}
FONT = "Arial"
W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
N = [0]


def _run(p, text, size=16, bold=False, color=INK, italic=False):
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold, f.italic, f.color.rgb = FONT, Pt(size), bold, italic, color
    return r


def textbox(slide, x, y, w, h, paras, size=16, color=INK, anchor=MSO_ANCHOR.TOP, align=None):
    """paras: list of str | (str, dict) ; '**x**' prefix not parsed - use dict(bold=True)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, para in enumerate(paras):
        text, opt = (para, {}) if isinstance(para, str) else para
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        p.space_after = Pt(opt.get("after", 6))
        lvl = opt.get("level", 0)
        bullet = opt.get("bullet", False)
        prefix = ("•  " if lvl == 0 else "–  ") if bullet else ""
        if lvl:
            p.level = lvl
        if prefix:
            _run(p, prefix, opt.get("size", size), False, ACCENT if lvl == 0 else MUTED)
        if "parts" in opt:
            for t, b in opt["parts"]:
                _run(p, t, opt.get("size", size), b, opt.get("color", color))
        else:
            _run(p, text, opt.get("size", size), opt.get("bold", False), opt.get("color", color),
                 opt.get("italic", False))
    return tb


def bullets(slide, x, y, w, h, items, size=17):
    paras = []
    for it in items:
        if isinstance(it, tuple):  # (bold head, rest)
            paras.append(("", {"bullet": True, "parts": [(it[0], True), (it[1], False)], "after": 9}))
        elif it.startswith("- "):
            paras.append((it[2:], {"bullet": True, "level": 1, "size": size - 2, "color": INK2, "after": 5}))
        else:
            paras.append((it, {"bullet": True, "after": 9}))
    return textbox(slide, x, y, w, h, paras, size)


def new_slide(title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAGE
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    if kicker:
        textbox(s, Inches(0.55), Inches(0.32), Inches(11), Inches(0.35),
                [(kicker.upper(), {"size": 12, "bold": True, "color": ACCENT})])
    textbox(s, Inches(0.55), Inches(0.62), Inches(12.2), Inches(0.8),
            [(title, {"size": 28, "bold": True})])
    N[0] += 1
    textbox(s, Inches(12.3), Inches(7.0), Inches(0.8), Inches(0.3),
            [(str(N[0]), {"size": 11, "color": MUTED})], align=PP_ALIGN.RIGHT)
    return s


def card(slide, x, y, w, h, fill=CARD, line=GRID):
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    r.adjustments[0] = 0.06
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    r.line.color.rgb = line
    r.line.width = Pt(0.75)
    r.shadow.inherit = False
    return r


def card_text(slide, x, y, w, h, head, body, head_color=INK, fill=CARD, size=14):
    card(slide, x, y, w, h, fill)
    textbox(slide, x + Inches(0.15), y + Inches(0.12), w - Inches(0.3), h - Inches(0.2),
            [(head, {"size": size + 2, "bold": True, "color": head_color, "after": 4}),
             (body, {"size": size, "color": INK2})])


def table(slide, x, y, w, rows, col_w=None, size=12, header_fill=SOFT, row_h=0.36, bold_col0=False):
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, x, y, w, Inches(row_h * nr))
    t = shp.table
    if col_w:
        tot = sum(col_w)
        for i, cw in enumerate(col_w):
            t.columns[i].width = Emu(int(w * cw / tot))
    for r in range(nr):
        t.rows[r].height = Inches(row_h)
        for c in range(nc):
            cell = t.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (CARD if r % 2 else PAGE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            val = rows[r][c]
            color = INK
            if isinstance(val, tuple):
                val, color = val
            _run(p, str(val), size, r == 0 or (bold_col0 and c == 0), color)
    return t


def picture(slide, path, x, y, max_w, max_h, border=True):
    im = Image.open(path)
    ar = im.width / im.height
    w, h = max_w, int(max_w / ar)
    if h > max_h:
        h, w = max_h, int(max_h * ar)
    pic = slide.shapes.add_picture(path, x + (max_w - w) // 2, y, w, h)
    if border:
        pic.line.color.rgb = GRID
        pic.line.width = Pt(0.75)
    return pic


def arrow(slide, x, y, w=Inches(0.35)):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, Inches(0.3))
    a.fill.solid()
    a.fill.fore_color.rgb = MUTED
    a.line.fill.background()


def robust_chart(slide, R, x, y, w, h):
    """Native clustered bar chart of detection rate per transform × detector."""
    cd = CategoryChartData()
    cd.categories = [t["label"] for t in R["transforms"]]
    for d in R["detectors"]:
        cd.add_series(d["label"], [R["tpr"][t["key"]][d["key"]]["rate"] for t in R["transforms"]])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size, ch.legend.font.name, ch.legend.font.color.rgb = Pt(11), FONT, INK2
    ca, va = ch.category_axis, ch.value_axis
    ca.reverse_order = True
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(12), FONT
    ca.format.line.color.rgb = GRID
    va.maximum_scale, va.minimum_scale, va.major_unit = 1.0, 0.0, 0.25
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0%", False
    va.tick_labels.font.size, va.tick_labels.font.color.rgb = Pt(10), MUTED
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = GRID
    va.format.line.fill.background()
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 60, -8
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = "0%", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.color.rgb, dl.font.name = Pt(9), INK2, FONT
    for i, ser in enumerate(plot.series):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = SERIES[i]
    return ch


IMG = lambda n: os.path.join(C.IMG, n)  # noqa: E731
STATUS_COLOR = {"Đạt": STATUS["good"], "Mô phỏng": RGBColor(0xB0, 0x7A, 0x00),
                "Thay thế": RGBColor(0xB0, 0x7A, 0x00), "Chưa": STATUS["critical"]}
