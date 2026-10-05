"""Builds report/BaoCao_PhishLens.docx"""
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import content as C  # noqa: E402

ACCENT = RGBColor(0x2A, 0x78, 0xD6)
MUTED = RGBColor(0x6B, 0x6A, 0x66)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.0)

st = doc.styles["Normal"]
st.font.name, st.font.size = "Arial", Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.2
for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
    h = doc.styles[f"Heading {lvl}"]
    h.font.name, h.font.size, h.font.bold = "Arial", Pt(size), True
    h.font.color.rgb = ACCENT if lvl < 3 else RGBColor(0x0B, 0x0B, 0x0B)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)


def para(text="", bold_prefix=None, italic=False, color=None, size=None, align=None):
    p = doc.add_paragraph()
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    r = p.add_run(text)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def bullet(text, bold_prefix=None, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    return p


def numbered(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Number")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)


def _shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def table(rows, widths=None, size=9.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            r = c.paragraphs[0].add_run(str(val))
            r.font.size = Pt(size)
            r.bold = i == 0
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            if i == 0:
                _shade(c, "CDE2FB")
            if widths:
                c.width = Cm(widths[j])
    doc.add_paragraph()
    return t


def figure(path, caption, width_cm=16):
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = para(caption, italic=True, color=MUTED, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_after = Pt(10)


IMG = lambda n: os.path.join(C.IMG, n)  # noqa: E731
R = C.robustness()

# ================================================================== cover
for _ in range(5):
    doc.add_paragraph()
p = para("BÁO CÁO", size=14, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(C.TITLE.upper())
r.bold, r.font.size, r.font.color.rgb = True, Pt(22), ACCENT
para(C.SUBTITLE, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(3):
    doc.add_paragraph()
para("Bài báo gốc: " + C.PAPER, italic=True, size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(4):
    doc.add_paragraph()
para(f"Người thực hiện: {C.AUTHOR}", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ================================================================== summary
doc.add_heading("Tóm tắt", 1)
para("Báo cáo trình bày ba phần việc. (1) Tái hiện phương pháp KD-BiLSTM của Eskandarian et al. "
     "(JISA 2026): chưng cất tri thức từ MobileBERT vào một BiLSTM có multi-head attention để phát "
     "hiện email lừa đảo, kể cả email do LLM viết lại. (2) Các phần bổ sung và nâng cao trong quá trình "
     "tái hiện. (3) Ứng dụng bài báo vào một hướng mới có tính thực tiễn hơn: PhishLens Edu, một lớp "
     "phân tích rủi ro email đa kênh, có giải thích bằng tiếng Việt, dành cho môi trường trường học, "
     "kèm thí nghiệm đánh giá độ bền trước email lừa đảo do LLM tạo và chuyển kênh.")
bullet("đã dựng lại đầy đủ dữ liệu, 4 baseline, teacher, KD-BiLSTM và quy trình 5 kịch bản × 5-fold; "
       "số tham số khớp bài báo (4,44M so với 4,5M), student nhanh hơn teacher ~34 lần trên CPU. "
       "Chưa so được số liệu tuyệt đối vì không có dữ liệu gốc.", "Tái hiện: ")
bullet("khi thông điệp lừa đảo bị chuyển vào ảnh, mã QR hay tệp đính kèm, mô hình chỉ đọc văn bản "
       "(như KD-BiLSTM) phát hiện được 0%; hệ thống đa kênh phục hồi 57–100%. LLM viết lại làm mô hình "
       "văn bản giảm từ 80% xuống 67%, hệ đa kênh giữ 80%.", "Kết quả chính: ")
bullet("dữ liệu tổng hợp, 'LLM viết lại' là bộ luật mô phỏng; con số thể hiện xu hướng.", "Lưu ý: ")

# ================================================================== 1 paper
doc.add_heading("1. Tóm tắt bài báo gốc", 1)
doc.add_heading("1.1 Bài toán", 2)
para("LLM giúp kẻ tấn công viết email lừa đảo trôi chảy, lịch sự, cá nhân hóa, khiến các bộ lọc dựa "
     "trên từ khóa và danh tiếng bị qua mặt. Các mô hình transformer lớn phát hiện tốt nhưng nặng, khó "
     "đặt trực tiếp ở mail gateway hay thiết bị đầu cuối. Mục tiêu của bài báo là một bộ phát hiện nhỏ, "
     "nhanh, chạy được trên CPU, nhưng vẫn bền trước email do LLM viết lại.")
doc.add_heading("1.2 Dữ liệu", 2)
para("Tác giả gộp 5 bộ dữ liệu thật (Cambridge, Nazario, Chakraborty, Phishing Pot, Enron; ~623.000 "
     "email) và dùng LLM thương mại (OpenAI, DeepSeek) viết lại email lừa đảo theo ba hành vi né tránh: "
     "paraphrase (diễn đạt lại), masking (làm mềm dấu hiệu lộ liễu) và personalization (chèn thông tin "
     "người nhận). Phân tích TF-IDF cho thấy email do LLM sinh gần email hợp lệ hơn email lừa đảo gốc. "
     "Tác giả còn dùng bộ câu hỏi chấm 'mức độ phishing' để lọc mẫu sinh kém và khử trùng lặp.")
doc.add_heading("1.3 Phương pháp", 2)
para("Các mô hình được nâng cấp dần: LSTM → BiLSTM → BiLSTM + single-head attention → BiLSTM + "
     "multi-head attention (H = 4, D = 64), đầu vào là Word2Vec 100 chiều. Bước cuối là knowledge "
     "distillation: teacher MobileBERT (25,3M tham số) được fine-tune rồi đóng băng; student BiLSTM + "
     "multi-head attention (~4,5M tham số) có embedding khởi tạo từ teacher, học theo hàm mất mát")
para("L = α·CE(y, z_S) + (1 − α)·τ²·KL(σ(z_T/τ) ‖ σ(z_S/τ)),  với α = 0,5 và τ = 2.",
     italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Đánh giá dùng 5 kịch bản (Orig-Orig, Gen-Gen, Orig-Gen, Gen-Orig, Mixture), 5-fold, chia "
     "train/val/test = 72/8/20.")
doc.add_heading("1.4 Kết quả", 2)
table([("Mô hình (Mixture)", "Tham số", "F1 (%)", "Thời gian test (s)")] + C.PAPER_MIXTURE,
      [7, 2.5, 2.2, 3.3])
para("KD-BiLSTM kém các transformer tốt nhất khoảng 1–2,5 điểm F1, nhưng nhanh hơn 5–19 lần và nhỏ "
     "hơn 20–800 lần. Mô hình chỉ học email LLM rồi gặp email thật (Gen-Orig) có recall rất thấp "
     "(BiLSTM 51%); multi-head attention và distillation cải thiện rõ khả năng tổng quát hóa. Dữ liệu và "
     "code của bài báo chưa công bố, một số chi tiết kiến trúc không được mô tả đầy đủ.")

# ================================================================== 2 reproduction
doc.add_heading("2. Tái hiện bài báo", 1)
doc.add_heading("2.1 Phạm vi và cách làm", 2)
para("Toàn bộ pipeline được viết lại bằng PyTorch/Transformers trong package phishkd, có script chạy "
     "thí nghiệm và một notebook Colab tự chứa. Do không có dữ liệu gốc, một corpus 'toy' được sinh từ "
     "template; phần sinh email bằng LLM được mô phỏng bằng bộ luật thực hiện đúng ba hành vi của bài "
     "báo (prompt cho LLM thật đã được chuẩn bị sẵn).")
doc.add_heading("2.2 Mức độ tái hiện", 2)
table([("Thành phần", "Bài báo", "Bản tái hiện", "Trạng thái")] + C.REPRO, [4.2, 4.0, 5.6, 2.0], size=9)
doc.add_heading("2.3 Số liệu đo được", 2)
para("Số tham số của LSTM, BiLSTM, teacher và student khớp với bài báo. Hai mô hình attention lệch do "
     "bài báo không mô tả đủ cách nối attention và pooling.")
table([("Mô hình", "Bài báo", "Tái hiện")] + C.PARAMS, [7, 4, 4])
para("Kịch bản chéo Orig→Gen (học email người viết, kiểm thử email 'LLM') trên dữ liệu toy, 1 fold:")
table([("Mô hình", "F1 bài báo (%)", "F1 tái hiện (%)")] + C.ORIG_GEN, [7, 4, 4])
para(f"Tốc độ suy luận: student {C.SPEED['ours'][0]} so với teacher {C.SPEED['ours'][1]} trên CPU "
     f"({C.SPEED['ours'][2]} nhanh hơn), cùng xu hướng với bài báo ({C.SPEED['paper'][0]} so với "
     f"{C.SPEED['paper'][1]}, {C.SPEED['paper'][2]} trên GPU). Ở các kịch bản cùng phân phối, dữ liệu toy "
     "quá dễ nên mọi mô hình đạt ~100%; các con số chỉ dùng để kiểm tra pipeline. Riêng kết quả teacher "
     "Orig→Gen thấp (51,66) cho thấy teacher nhạy với độ lệch phân phối trên tập train nhỏ.")
doc.add_heading("2.4 Vấn đề gặp phải và cách khắc phục", 2)
table([("Hiện tượng", "Nguyên nhân", "Khắc phục")] + C.FIXES, [5, 5.5, 5.5])
doc.add_heading("2.5 Những chỗ tự quyết định", 2)
for t in ["Cách chia fold cho kịch bản chéo (Orig-Gen, Gen-Orig): K-fold song song trên hai miền.",
          "Chọn checkpoint theo weighted-F1 trên tập validation 8%.",
          "Pooling sau attention: masked average pooling; student có residual + LayerNorm.",
          "Logit của teacher được tính một lần rồi dùng lại (tương đương Algorithm 1 vì teacher đóng băng).",
          "Ngưỡng khử trùng lặp, ngưỡng lọc theo mức phishing, trọng số câu hỏi theo AUC."]:
    bullet(t)

# ================================================================== 3 extras
doc.add_heading("3. Các phần bổ sung và nâng cao", 1)
for t in C.EXTRAS:
    bullet(t)
para("Trong đó, hai điểm có ý nghĩa phương pháp luận nhất là chia template held-out (tránh kết quả "
     "~100% giả tạo khi train và test dùng chung cách viết) và che URL trước mô hình văn bản. Phiên bản "
     "đầu của mô hình văn bản đã báo nhầm cả email thật có link chính thức của trường, vì trong dữ liệu "
     "tổng hợp chỉ email lừa đảo mới chứa tên miền: mô hình học 'có tên miền = lừa đảo'. Sau khi che URL "
     "(liên kết đã có bộ phân tích riêng) và bổ sung email hợp lệ 'khó', hiện tượng này được khắc phục.")

# ================================================================== 4 application
doc.add_heading("4. Ứng dụng bài báo: PhishLens Edu", 1)
doc.add_heading("4.1 Động cơ", 2)
para("Khi đem bài báo vào thực tế, ba vấn đề nổi lên. Thứ nhất, các trường đã có bộ lọc của Gmail/"
     "Microsoft 365, nên thêm một bộ phân loại văn bản ít giá trị. Thứ hai, email thật có ảnh và tệp: kẻ "
     "tấn công giấu thông điệp vào ảnh, mã QR, PDF, HTML, ZIP có mật khẩu, mà mô hình chỉ đọc thân thư "
     "không thấy. Thứ ba, đầu ra có/không không giúp người dùng hiểu vì sao email nguy hiểm.")
para("PhishLens Edu không thay thế Gmail/Defender mà là lớp “ý kiến thứ hai có giải thích” cho những "
     "email đã lọt vào hộp thư: người dùng chuyển tiếp email nghi ngờ (hoặc tải tệp .eml lên) và nhận "
     "về mức độ rủi ro, lý do và hướng dẫn. Toàn bộ xử lý chạy cục bộ; tệp đính kèm chỉ được phân tích "
     "tĩnh, không bao giờ được mở hay thực thi.")
doc.add_heading("4.2 Kiến trúc", 2)
for h, d in [("Tách email: ", "đọc .eml (MIME) thành header, văn bản/HTML, liên kết (chữ hiển thị và "
              "địa chỉ thật), ảnh, tệp đính kèm."),
             ("Phân tích đa kênh: ", "5 bộ phân tích độc lập, mỗi dấu hiệu có độ mạnh và lời giải thích."),
             ("Kết hợp bằng chứng: ", "gộp theo noisy-OR thành điểm từng kênh và điểm tổng 0–100, chia 4 "
              "mức (An toàn / Cần chú ý / Nguy hiểm / Rất nguy hiểm); suy ra loại tấn công và kỹ thuật né tránh."),
             ("Giải thích & hướng dẫn: ", "danh sách lý do kèm bài học, tô sáng cụm từ đáng ngờ, lời khuyên "
              "theo mức độ và loại tấn công.")]:
    numbered(d, h)
table([("Kênh", "Phân tích")] + C.CHANNELS, [3, 13])
doc.add_heading("4.3 Kế thừa bài báo", 2)
para("Kênh nội dung dùng chính KD-BiLSTM của bài báo, huấn luyện lại trên email trường học tiếng Việt "
     "và email tiếng Anh (có cả bản viết lại kiểu LLM). Trọng số attention được ánh xạ về chữ gốc có dấu "
     "để hiển thị những từ mô hình chú ý. Mô hình chỉ là một trong năm kênh: khi văn bản bị giấu vào "
     "ảnh hay tệp, các kênh khác bù lại; ngược lại, với email không có link/tệp (mạo danh nhờ chuyển "
     "tiền), mô hình nội dung và các dấu hiệu tâm lý là nguồn tín hiệu chính.")
doc.add_heading("4.4 Giao diện và demo", 2)
para("Giao diện web (chạy cục bộ) có bốn tab: Phân tích email, Luyện tập nhận diện, Đánh giá độ bền, "
     "Giới thiệu. Bộ 12 email mẫu bao phủ email thật, lừa đảo cổ điển, email viết lại kiểu LLM, ảnh + QR, "
     "HTML đăng nhập giả, macro Word, ZIP có mật khẩu chứa .pdf.exe, mạo danh thầy nhờ mua thẻ cào, PDF "
     "có JavaScript. Mọi tệp 'độc hại' đều là tệp giả vô hại.")
table([("Email mẫu", "Điểm", "Mức")] + C.SAMPLES, [10, 2, 4])
figure(IMG("s04_qr_m365_top.png"), "Hình 1. Email chỉ có ảnh + mã QR: điểm rủi ro, lý do và mức rủi ro theo kênh.")
figure(IMG("c_qr_figure.png"), "Hình 2. OCR đọc chữ tiếng Việt trong ảnh và giải mã mã QR ra liên kết giả mạo.", 8)
figure(IMG("c_zip_att.png"), "Hình 3. ZIP có mật khẩu: vẫn đọc được tên tệp bên trong và phát hiện đuôi kép .pdf.exe.")
figure(IMG("s08_the_cao_top.png"), "Hình 4. Mạo danh thầy nhờ mua thẻ cào: không link, không tệp, mô hình văn "
       "bản chỉ cho 4% nhưng hệ đa kênh chấm 78/100.")
figure(IMG("quiz_a.png"), "Hình 5. Chế độ luyện tập: người dùng đoán trước, sau đó xem đáp án và phân tích.")

# ================================================================== 5 robustness
doc.add_heading("5. Đánh giá độ bền", 1)
doc.add_heading("5.1 Thiết kế thí nghiệm", 2)
para("Câu hỏi nghiên cứu: (RQ1) bộ phát hiện chỉ dựa trên văn bản suy giảm thế nào khi email lừa đảo "
     "được LLM viết lại hoặc chuyển sang ảnh, mã QR, tệp đính kèm? (RQ2) kết hợp bằng chứng đa kênh có "
     "khôi phục khả năng phát hiện mà không tăng báo nhầm không?")
para(f"Mỗi kịch bản dùng {R['n_phish']} email lừa đảo và {R['n_legit']} email hợp lệ lấy từ template "
     "held-out (chưa dùng để huấn luyện). Email hợp lệ được biến đổi giống hệt để đo báo nhầm; header "
     "người gửi giữ trung tính để chỉ còn 'cách mang thông điệp' thay đổi. Ngưỡng: điểm ≥ 50 cho hệ đa "
     "kênh, P ≥ 0,5 cho mô hình chỉ đọc văn bản.")
table([("Kịch bản", "Mô tả")] + [(t["label"], t["desc"]) for t in R["transforms"]], [4, 12])
doc.add_heading("5.2 Kết quả", 2)
fmt = lambda v: f"{100 * v:.0f}%"  # noqa: E731
dk = [d["key"] for d in R["detectors"]]
rows = [("Kịch bản", "VB TF-IDF", "VB KD", "ĐK TF-IDF", "ĐK KD", "Báo nhầm ĐK TF-IDF", "Báo nhầm ĐK KD")]
for t in R["transforms"]:
    k = t["key"]
    rows.append((t["label"], *[fmt(R["tpr"][k][d]["rate"]) for d in dk],
                 fmt(R["fpr"][k]["full_linear"]["rate"]), fmt(R["fpr"][k]["full_kd"]["rate"])))
table(rows, [3.4, 1.9, 1.9, 1.9, 1.9, 2.5, 2.5], size=9)
para("VB = chỉ đọc văn bản thân thư; ĐK = PhishLens đa kênh; KD = KD-BiLSTM của bài báo.",
     italic=True, color=MUTED, size=9.5)
figure(IMG("c_robust_chart.png"), "Hình 6. Tỉ lệ phát hiện email lừa đảo theo kịch bản và bộ phát hiện.")
doc.add_heading("5.3 Phân tích", 2)
for h, d in [
    ("Chuyển kênh vô hiệu hóa mô hình chỉ đọc văn bản. ",
     "Khi thông điệp nằm trong ảnh, mã QR, tệp Word hay HTML, mô hình chỉ đọc thân thư phát hiện 0%; "
     "hệ đa kênh phục hồi 57–100% nhờ OCR, giải mã QR và trích xuất văn bản từ tệp."),
    ("LLM viết lại làm giảm mô hình văn bản. ",
     "TF-IDF giảm từ 80% xuống 67%, KD-BiLSTM từ 67% xuống 60%; hệ đa kênh giữ 80%."),
    ("BEC là loại khó nhất. ",
     "Mạo danh nhờ chuyển tiền/mua thẻ cào không có link hay tệp; khi header người gửi bị giữ trung tính, "
     "hệ đa kênh vẫn bỏ sót 50%. Trong thực tế, tín hiệu mạnh nhất là người gửi mạo danh (mẫu demo 78/100)."),
    ("Giới hạn của mô hình bài báo với tiếng Việt. ",
     "Trên email hợp lệ chưa thấy, KD-BiLSTM báo nhầm 14% so với 1% của TF-IDF (trong hệ đa kênh); "
     "nguyên nhân khả dĩ là tokenizer MobileBERT tiếng Anh xóa dấu. Cần teacher đa ngôn ngữ (PhoBERT, "
     "XLM-R) hoặc hiệu chỉnh ngưỡng."),
    ("Tách bạch các kênh giúp hệ thống bền hơn. ",
     "Bài học 'học đường tắt' (mục 3) cho thấy mô hình văn bản dễ học mẹo từ dữ liệu; để liên kết, người "
     "gửi, ảnh và tệp cho các bộ phân tích chuyên biệt giúp giảm phụ thuộc vào một mô hình duy nhất."),
]:
    bullet(d, h)

# ================================================================== 6 limits
doc.add_heading("6. Hạn chế và hướng phát triển", 1)
doc.add_heading("6.1 Hạn chế", 2)
for t in ["Dữ liệu tổng hợp; chưa có email lừa đảo thật của trường.",
          "'LLM viết lại' là bộ luật mô phỏng, chưa dùng LLM thật.",
          "Chưa tái hiện Table 7–8 của bài báo (ModernBERT, DeBERTaV3, T5, Qwen, Phi-4, LR, XGBoost, "
          "fastText, TinyBERT, DistilBERT).",
          "Tokenizer MobileBERT không giữ dấu tiếng Việt.",
          "Chưa kiểm tra danh tiếng URL trực tuyến, chưa chạy tệp trong sandbox; OCR trên CPU mất vài giây mỗi ảnh."]:
    bullet(t)
doc.add_heading("6.2 Hướng phát triển", 2)
for t in ["Thu thập email lừa đảo thật (ẩn danh hóa) qua hộp thư báo cáo của Phòng CNTT.",
          "Sinh biến thể bằng LLM chạy cục bộ (đã có chỗ cắm Ollama) để tạo dữ liệu đa dạng hơn.",
          "Chưng cất từ teacher đa ngôn ngữ (PhoBERT, XLM-R) cho tiếng Việt.",
          "Khảo sát người dùng: bài kiểm tra nhận diện trước và sau 2 tuần dùng chế độ luyện tập (RQ3).",
          "Tích hợp: add-in Outlook/Gmail hoặc hộp thư “chuyển tiếp để kiểm tra”."]:
    bullet(t)

# ================================================================== 7 conclusion
doc.add_heading("7. Kết luận", 1)
para("Phương pháp KD-BiLSTM của bài báo đã được tái hiện đầy đủ về mặt kỹ thuật; số tham số và các xu "
     "hướng (chéo phân phối làm giảm điểm, multi-head và distillation giúp tổng quát hóa, student nhanh "
     "hơn teacher nhiều lần) khớp bài báo, dù chưa so được số liệu tuyệt đối do thiếu dữ liệu gốc. Khi "
     "ứng dụng vào thực tế, mô hình được đặt làm kênh nội dung trong PhishLens Edu, một hệ thống phân tích "
     "đa kênh có giải thích cho trường học. Thí nghiệm độ bền cho thấy chỉ đọc văn bản là không đủ trước "
     "các kỹ thuật chuyển kênh (0%), trong khi kết hợp bằng chứng đa kênh phục hồi 57–100% khả năng phát "
     "hiện. Bước tiếp theo quan trọng nhất là kiểm chứng trên email thật và với LLM thật.")

# ================================================================== appendix
doc.add_heading("Phụ lục: mã nguồn và cách chạy", 1)
table([("Thành phần", "Đường dẫn"),
       ("Tái hiện bài báo", "phishkd/, scripts/run_experiments.py, phishing_kd_bilstm.ipynb"),
       ("PhishLens Edu", "phishlens/ (parser, sender_links, content, media, engine, app, web/)"),
       ("Dữ liệu & mô hình", "scripts/build_samples.py, scripts/train_text_models.py --kd"),
       ("Thí nghiệm độ bền", "scripts/robustness_eval.py → results/robustness.json"),
       ("Chạy giao diện", "uvicorn phishlens.app:app --host 127.0.0.1 --port 8765")], [4.5, 11.5])

out = os.path.join(C.ROOT, "report", "BaoCao_PhishLens.docx")
doc.save(out)
print("saved", out)
